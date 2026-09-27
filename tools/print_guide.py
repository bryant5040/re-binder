"""Makes a printable Word copy of the owner's guide, the Read Me First, for teaching day.

The guide has one source: plugin/skills/binder-setup/references/read-me-first.md, which
"set up my binder" puts in his binder word for word. This only lays that same text out on a page,
so the two can never drift apart. Run it again after any change to the guide.

Usage: python tools/print_guide.py ["<output .docx>"]
Default output: dist/printable/Read Me First.docx (dist/ is never committed).
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "plugin" / "skills" / "binder-setup" / "references" / "read-me-first.md"
DEFAULT_OUT = ROOT / "dist" / "printable" / "Read Me First.docx"


def blocks(text):
    """The guide's sections: a title line, then its lines, with wrapped lines joined back up."""
    for chunk in re.split(r"\n\s*\n", text.strip()):
        lines = []
        for line in chunk.splitlines():
            if line.startswith("  ") and lines:
                lines[-1] += " " + line.strip()  # a wrapped line belongs to the one above
            else:
                lines.append(line.strip())
        yield lines


def main(argv):
    out = Path(argv[1]) if len(argv) > 1 else DEFAULT_OUT
    doc = Document()
    for section in doc.sections:  # one page, to keep in the truck
        section.top_margin = section.bottom_margin = Inches(0.6)
        section.left_margin = section.right_margin = Inches(0.7)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(1)
    for name, size in (("Title", 20), ("Heading 2", 12.5)):
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].paragraph_format.space_before = Pt(7)
        doc.styles[name].paragraph_format.space_after = Pt(2)
    sections = list(blocks(SOURCE.read_text(encoding="utf-8")))
    doc.add_heading(sections[0][0], level=0)
    for lines in sections[1:]:
        doc.add_heading(lines[0], level=2)
        for line in lines[1:]:
            if line.startswith("- "):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif re.match(r"^\d+\.\s", line):
                doc.add_paragraph(re.sub(r"^\d+\.\s", "", line), style="List Number")
            else:
                doc.add_paragraph(line)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main(sys.argv)
